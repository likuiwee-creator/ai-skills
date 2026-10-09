# -*- coding: utf-8 -*-
"""
携程(景点玩乐代理)账单明细提取脚本（按月复用）
用法:
  python extract_ctrip_bill.py [输入xlsx] [输出xlsx可选]

逻辑:
  1. 取「订单明细」sheet 作为主表
  2. 去除顶部 SKIP_ROWS 行（默认2行：标题行 + To: 抬头行）
  3. 表头重命名：携程订单号 -> OTA订单号、结算价金额 -> 打款金额
  4. 裁掉表头为空的尾部多余列，输出为单 sheet 表格
"""
import openpyxl
import sys

SRC = sys.argv[1] if len(sys.argv) > 1 else r"D:/360极速浏览器X下载/景点玩乐代理-2026.07.01-2026.07.31(1201448477) (1).xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else SRC[: SRC.rfind(".")] + "_账单明细.xlsx"

SKIP_ROWS = 2
# 优先按这些关键字匹配目标 sheet（兼容「订单明细」「账单明细」两种叫法）
TARGET_KWS = ["订单明细", "账单明细", "明细"]
RENAME = {"携程订单号": "OTA订单号", "结算价金额": "打款金额"}


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

    target = None
    for kw in TARGET_KWS:
        for n in wb.sheetnames:
            if kw in n:
                target = n
                break
        if target:
            break
    if target is None:
        target = wb.sheetnames[0]
        print("未找到含明细关键字的sheet，已回退使用第一个sheet: %s" % target)

    ws = wb[target]
    rows = list(ws.iter_rows(values_only=True))
    data = rows[SKIP_ROWS:]  # data[0] 现在是真正的表头

    header = list(data[0])
    # 找到最后一个非空表头列，裁掉尾部全空列
    last = max(i for i, h in enumerate(header) if h is not None)
    header = header[: last + 1]
    header = [RENAME.get(str(h).strip(), h) for h in header]

    out_rows = [header]
    for r in data[1:]:
        row = list(r)
        if len(row) > last + 1:
            row = row[: last + 1]
        elif len(row) < last + 1:
            row = row + [None] * (last + 1 - len(row))
        out_rows.append(row)

    out = openpyxl.Workbook()
    o = out.active
    o.title = "账单明细"
    for r in out_rows:
        o.append(r)
    out.save(OUT)

    print("已保存:", OUT)
    print("来源sheet:", target, " 跳过顶部行数:", SKIP_ROWS)
    print("输出行数(含表头):", len(out_rows), " 输出列数:", len(header))
    print("新表头前几列:", header[:4], "...", header[-2:])


if __name__ == "__main__":
    main()
